"""Écriture du rapport Excel (feuilles Résultats, Synthèse, Config)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .utils import flatten_dict, to_cell_value

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")


def build_summary(results: pd.DataFrame, metric_names: list[str]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for name in metric_names:
        scores = pd.to_numeric(results.get(f"{name}_score"), errors="coerce")
        passed = results.get(f"{name}_passed")
        evaluated = passed.dropna() if passed is not None else pd.Series(dtype=object)
        rows.append(
            {
                "métrique": name,
                "cas évalués": int(evaluated.size),
                "score moyen": round(scores.mean(), 4) if scores.notna().any() else None,
                "score min": round(scores.min(), 4) if scores.notna().any() else None,
                "taux de réussite": round(evaluated.astype(bool).mean(), 4) if evaluated.size else None,
            }
        )
    return pd.DataFrame(rows)


def build_run_stats(results: pd.DataFrame) -> pd.DataFrame:
    latency = pd.to_numeric(results.get("latency_s"), errors="coerce")
    errors = results["error"].notna().sum() if "error" in results else 0
    stats = {
        "cas de test": len(results),
        "appels en erreur": int(errors),
        "latence moyenne (s)": round(latency.mean(), 3) if latency.notna().any() else None,
        "latence p95 (s)": round(latency.quantile(0.95), 3) if latency.notna().any() else None,
    }
    return pd.DataFrame({"indicateur": list(stats), "valeur": list(stats.values())})


def _style_sheet(ws, wrap_columns: set[str] | None = None, max_width: int = 60) -> None:
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, column_cells in enumerate(ws.columns, start=1):
        header = str(column_cells[0].value or "")
        sample = [len(str(c.value)) for c in column_cells[:200] if c.value is not None]
        width = min(max([len(header), *sample, 8]) + 2, max_width)
        ws.column_dimensions[get_column_letter(index)].width = width
        if wrap_columns and header in wrap_columns:
            for cell in column_cells[1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def _color_pass_fail(ws, headers: list[str]) -> None:
    last_row = ws.max_row
    if last_row < 2:
        return
    for index, header in enumerate(headers, start=1):
        if header.endswith("_passed"):
            letter = get_column_letter(index)
            cell_range = f"{letter}2:{letter}{last_row}"
            ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=["TRUE"], fill=GREEN))
            ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=["FALSE"], fill=RED))


def write_report(
    path: str | Path,
    results: pd.DataFrame,
    metric_names: list[str],
    config_info: dict[str, Any],
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    results = results.map(to_cell_value)
    summary = build_summary(results, metric_names)
    stats = build_run_stats(results)

    config_rows = flatten_dict(config_info)
    config_rows["date d'exécution"] = datetime.now().isoformat(timespec="seconds")
    config_df = pd.DataFrame({"paramètre": list(config_rows), "valeur": [to_cell_value(v) for v in config_rows.values()]})

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        results.to_excel(writer, sheet_name="Résultats", index=False)
        summary.to_excel(writer, sheet_name="Synthèse", index=False)
        stats.to_excel(writer, sheet_name="Synthèse", index=False, startrow=len(summary) + 3)
        config_df.to_excel(writer, sheet_name="Config", index=False)

        ws = writer.sheets["Résultats"]
        headers = list(results.columns)
        wrap = {h for h in headers if h in {"input", "answer", "actual_output"} or h.endswith(("_reason", "[json]"))}
        _style_sheet(ws, wrap_columns=wrap)
        _color_pass_fail(ws, headers)
        _style_sheet(writer.sheets["Synthèse"])
        _style_sheet(writer.sheets["Config"], max_width=100)
    return path


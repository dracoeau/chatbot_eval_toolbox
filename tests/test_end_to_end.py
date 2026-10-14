import json

import httpx
import openpyxl
import pytest

from chatbot_eval.cli import main
from chatbot_eval.config import load_config
from chatbot_eval.runner import run_evaluation

ANSWERS = {
    "Bonjour": "Bonjour ! Comment puis-je vous aider ?",
    "Quel délai ?": "Comptez environ une semaine.",
}


def make_transport(calls, fail_first=False):
    state = {"failed": False}

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        if fail_first and not state["failed"]:
            state["failed"] = True
            return httpx.Response(503, text="indisponible")
        body = json.loads(request.content)
        question = body["input"]["input"]
        return httpx.Response(
            200,
            json={"output": ANSWERS[question], "metadata": {"run_id": f"r-{question[:3]}", "sources": ["doc1"]}},
        )

    return httpx.MockTransport(handler)


def test_full_run(config_factory):
    config = config_factory(
        metrics=[
            {"type": "exact_match", "params": {"strip_punctuation": True}},
            {"type": "contains", "name": "mots_cles", "params": {"substrings_column": "keywords"}},
        ]
    )
    calls = []
    results, path = run_evaluation(config, transport=make_transport(calls, fail_first=True))

    # 1 échec 503 + retry + 1 autre appel
    assert len(calls) == 3
    request = calls[-1]
    assert str(request.url) == "https://bot.example.com/usecases/uc1/invoke"
    assert request.headers["Authorization"] == "Bearer secret-token"

    payloads = {json.loads(c.content)["input"]["input"]: json.loads(c.content) for c in calls}
    assert payloads["Bonjour"]["input"]["filters"] == {"domaine": "general", "top_k": 3}
    assert payloads["Quel délai ?"]["input"]["chat_history"][0]["role"] == "user"

    assert results["exact_match_passed"].tolist() == [True, False]
    assert results["mots_cles_passed"].tolist() == [True, False]
    assert results["chatbot.metadata.run_id"].tolist() == ["r-Bon", "r-Que"]

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Résultats", "Synthèse", "Config"]
    headers = [c.value for c in wb["Résultats"][1]]
    for expected in ["id", "Input", "Answer", "chat_history[json]", "kwargs.filters[json]", "actual_output",
                     "latency_s", "chatbot.output", "chatbot.metadata.sources", "exact_match_score",
                     "mots_cles_reason"]:
        assert expected in headers
    config_values = [str(c.value) for row in wb["Config"].iter_rows() for c in row]
    assert not any("secret-token" in v for v in config_values)


def test_http_error_does_not_stop_run(config_factory):
    config = config_factory(metrics=[{"type": "exact_match"}])
    transport = httpx.MockTransport(lambda request: httpx.Response(401, text="unauthorized"))
    results, _ = run_evaluation(config, transport=transport)
    assert results["status_code"].tolist() == [401, 401]
    assert results["error"].str.contains("HTTP 401").all()
    assert results["exact_match_passed"].isna().all()


def test_custom_metric_module(config_factory, tmp_path):
    module = tmp_path / "my_metrics.py"
    module.write_text(
        "from chatbot_eval import Metric, MetricResult, register_metric\n"
        "@register_metric('always_ok')\n"
        "class AlwaysOk(Metric):\n"
        "    def evaluate(self, sample):\n"
        "        return MetricResult(score=1.0, passed=True, reason='ok')\n"
    )
    config = config_factory(metrics=[{"type": "always_ok"}], custom_metrics_modules=["my_metrics.py"])
    results, _ = run_evaluation(config, transport=make_transport([]), config_dir=tmp_path)
    assert results["always_ok_passed"].all()


def test_env_interpolation(tmp_path, monkeypatch, dataset_path):
    monkeypatch.setenv("BOT_TOKEN", "from-env")
    path = tmp_path / "c.json"
    path.write_text(json.dumps({
        "chatbot": {"base_url": "x", "use_case_id": "y", "token": "${BOT_TOKEN}"},
        "dataset_path": str(dataset_path),
        "output_path": "out.xlsx",
    }))
    config = load_config(path)
    assert config.chatbot.token == "from-env"
    assert config.output_path == tmp_path / "out.xlsx"


def test_cli_invalid_config(tmp_path):
    path = tmp_path / "c.json"
    path.write_text("{}")
    assert main(["--config", str(path)]) == 2
